import os
import asyncio

from docx import Document


GENERATED_DIR = "generated"

os.makedirs(
    GENERATED_DIR,
    exist_ok=True
)



async def save_summary_file(
    job_id: str,
    filename: str,
    summary: str
):

    def _create_docx():
        document = Document()

        document.add_heading(
            "AI Generated Summary",
            level=1
        )

        document.add_paragraph(
            f"Original File: {filename}"
        )

        document.add_paragraph("")

        document.add_heading(
            "Summary",
            level=2
        )

        document.add_paragraph(summary)

        file_path = (
            f"{GENERATED_DIR}/{job_id}.docx"
        )

        document.save(file_path)

        return file_path

    return await asyncio.to_thread(_create_docx)
